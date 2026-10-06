"""Stage 1 entry point: turn any supported file into SourceDocument records.

Supported inputs
- PDF                 -> render pages -> Tesseract OCR -> field extraction
- PNG / JPG / TIFF    -> Tesseract OCR -> field extraction
- DOCX                -> paragraph + table text -> field extraction
- CSV / XLSX          -> one source document per row (headers become labels)

Every document keeps its source path, page images, OCR text, OCR confidence,
the extraction method used, and per-field confidence. Labels and values are
never renamed or normalized here.
"""

from __future__ import annotations

from pathlib import Path

from ingestion.extraction.dispatcher import parse_document
from ingestion.schemas.contracts import SourceDocument, source_document_from_extraction
from ingestion.schemas.invoice_schema import DocumentExtraction
from utils.config import LOW_CONFIDENCE_THRESHOLD

PDF_TYPES = {".pdf"}
IMAGE_TYPES = {".png", ".jpg", ".jpeg", ".tif", ".tiff"}
SHEET_TYPES = {".csv", ".xlsx"}
DOCX_TYPES = {".docx"}
SUPPORTED_TYPES = PDF_TYPES | IMAGE_TYPES | SHEET_TYPES | DOCX_TYPES


def _docx_text(path: Path) -> str:
    from docx import Document

    document = Document(str(path))
    lines = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            lines.append("  ".join(cell.text for cell in row.cells))
    return "\n".join(lines)


def _build(
    document_id: str,
    path: Path,
    extraction: dict,
    *,
    method: str,
    ocr_text: str | None,
    ocr_confidence: float | None,
    page_images: list[str],
    row_number: int | None = None,
) -> SourceDocument:
    validated = DocumentExtraction(**extraction)
    source = source_document_from_extraction(document_id, str(path), validated)
    confidences = [field.confidence for field in source.fields if field.confidence is not None]
    mean_confidence = round(sum(confidences) / len(confidences), 4) if confidences else 0.0
    return source.model_copy(
        update={
            "ocr_text": ocr_text,
            "pipeline_version": "stage-1/1.1",
            "source_name": path.name,
            "source_row": row_number,
            "page_images": page_images,
            "extraction_method": method,
            "ocr_confidence": ocr_confidence,
            "extraction_confidence": mean_confidence,
            "low_confidence": mean_confidence < LOW_CONFIDENCE_THRESHOLD,
        }
    )


def ingest_file(path: Path) -> list[SourceDocument]:
    """Extract one file into one or more source documents."""
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_TYPES:
        raise ValueError(f"Unsupported file type: {suffix or 'unknown'}")

    if suffix in SHEET_TYPES:
        from ingestion.tabular.tabular_reader import extract_rows

        return [
            _build(
                f"{path.stem}-row-{index:04d}",
                path,
                extraction,
                method="spreadsheet",
                ocr_text=None,
                ocr_confidence=None,
                page_images=[],
                row_number=index,
            )
            for index, extraction in enumerate(extract_rows(path), start=1)
        ]

    if suffix in DOCX_TYPES:
        text = _docx_text(path)
        extraction, method = parse_document(text)
        return [
            _build(path.stem, path, extraction, method=method, ocr_text=text,
                   ocr_confidence=None, page_images=[])
        ]

    from ingestion.ocr.tesseract_engine import extract_text_with_confidence

    if suffix in PDF_TYPES:
        from ingestion.pdf.pdf_to_image import pdf_to_images

        images = pdf_to_images(path)
    else:
        images = [path]
    text, ocr_confidence = extract_text_with_confidence(images)
    extraction, method = parse_document(text, ocr_confidence)
    return [
        _build(
            path.stem,
            path,
            extraction,
            method=method,
            ocr_text=text,
            ocr_confidence=ocr_confidence,
            page_images=[str(image) for image in images],
        )
    ]
